"""Export-only workbook metadata must not alter existing optical contracts."""
from dataclasses import fields, replace
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string

from glass_workbench.catalogs import cdgm, schott
from glass_workbench.contracts import AgfMetadata, Glass
from glass_workbench.optics.dispersion import air_to_vacuum, catalog_index

ROOT = Path(__file__).resolve().parents[1]


def put(rows, row, column, value):
    rows[row - 1][column_index_from_string(column) - 1] = value


@pytest.fixture(params=["CDGM", "SCHOTT"])
def source(request):
    supplier = request.param
    rows = [[None] * 240 for _ in range(3 if supplier == "CDGM" else 5)]
    if supplier == "CDGM":
        header_row, data_row, adapter = 1, 3, cdgm
        headers = {"A": "Glass", "B": "Code", "N": "nd", "X": "νd",
                   "AB": "K1", "AC": "L1", "AD": "K2", "AE": "L2", "AF": "K3", "AG": "L3",
                   "AH": "A0", "AM": "A5", "AO": "D0", "AT": "λTK(μm)",
                   "BZ": "α-30/70℃(10-7/K)", "CA": "α100/300℃(10-7/K)",
                   "CK": "ρ(g/cm3)", "BJ": "ΔPg,F", "EC": "τ(10mm)",
                   "HV": "Melt Frequency", "HW": "Relative Price"}
        values = {"A": "TEST", "N": 1.5, "X": 60, "BZ": 114, "CA": 143,
                  "CK": 4.11, "BJ": 0.0229, "AY": 0.55, "C": 1.51, "D": 1.49,
                  "AB": 1.0, "AD": 0.2, "AF": 0.8, "AC": 0.01, "AE": 0.02, "AG": 100.0}
        put(rows, 2, "C", 400.0)
        put(rows, 2, "D", 1529.6)
        columns = {"cte1": "BZ", "cte2": "CA", "density": "CK", "dpgf": "BJ"}
    else:
        header_row, data_row, adapter = 4, 5, schott
        put(rows, 1, "A", "SCHOTT, Optical Glass")
        headers = {"A": "Glass", "B": "nd", "D": "vd", "G": "B1", "H": "B2", "I": "B3",
                   "J": "C1", "K": "C2", "L": "C3", "M": "D0", "R": "lambda",
                   "DF": "alpha -30/70", "DG": "alpha 20/300", "CZ": "Density",
                   "FA": "  DP g,F", "FC": "Glascode", "DM": "  n2325.4", "DN": "  n1970.1"}
        values = {"A": "TEST", "B": 1.5, "D": 60, "DF": 7.4, "DG": 8.5,
                  "CZ": 5.41, "FA": 0.0033, "EP": 0.5982, "DM": 1.49, "DN": 1.5,
                  "G": 1.0, "H": 0.2, "I": 0.8, "J": 0.01, "K": 0.02, "L": 100.0}
        columns = {"cte1": "DF", "cte2": "DG", "density": "CZ", "dpgf": "FA"}
    for column, value in headers.items():
        put(rows, header_row, column, value)
    for column, value in values.items():
        put(rows, data_row, column, value)
    return supplier, rows, adapter, header_row, data_row, columns


def parse(source):
    return source[2].parse(source[1], "test.xlsx", "first", "digest")[0][0]


def test_contract_is_appended_and_optional():
    glass = Glass("TEST", "TEST", "", 1.5, 60)
    assert glass.agf_metadata is None
    assert fields(Glass)[-1].name == "agf_metadata"
    assert AgfMetadata().cte2 is None
    assert AgfMetadata().provenance is not AgfMetadata().provenance
    # Frozen Glass was hashable before the export-only field was added.
    hash(replace(glass, agf_metadata=AgfMetadata(provenance={"cte1": "source"})))


def test_verified_columns_units_and_direct_deviation(source):
    glass = parse(source)
    metadata = glass.agf_metadata
    assert metadata.cte1 == glass.cte
    if source[0] == "CDGM":
        assert metadata.cte1 == pytest.approx(11.4e-6)
        assert metadata.cte2 == pytest.approx(14.3e-6)
        assert metadata.density == 4.11
        assert metadata.dpgf == 0.0229
        assert metadata.cte2_temperature_range == (100, 300)
        assert metadata.wavelength_range_air == pytest.approx((0.4, 1.5296))
    else:
        assert metadata.cte1 == pytest.approx(7.4e-6)
        assert metadata.cte2 == pytest.approx(8.5e-6)
        assert metadata.density == 5.41
        assert metadata.dpgf == 0.0033
        assert metadata.cte2_temperature_range == (20, 300)
        assert metadata.wavelength_range_air == (1.97009, 2.32542)
        assert "not 100/300" in metadata.provenance["cte2"]
    assert metadata.dpgf != glass.pgf
    for key, column in source[5].items():
        assert f"first!{column}{source[4]}" in metadata.provenance[key]
    assert tuple(air_to_vacuum(w) for w in metadata.wavelength_range_air) == glass.wavelength_range


@pytest.mark.parametrize("missing", [None, "", "unknown", float("nan"), float("inf"), True])
def test_missing_or_invalid_values_never_become_zero(source, missing):
    for column in source[5].values():
        put(source[1], source[4], column, missing)
    metadata = parse(source).agf_metadata
    assert (metadata.cte1, metadata.cte2, metadata.density, metadata.dpgf) == (None,) * 4


def test_direct_zero_is_not_treated_as_missing(source):
    put(source[1], source[4], source[5]["dpgf"], 0)
    assert parse(source).agf_metadata.dpgf == 0.0


def test_optional_headers_do_not_change_legacy_fields_or_calculations(source):
    before = parse(source)
    for key in ("cte2", "density", "dpgf"):
        put(source[1], source[3], source[5][key], "unrecognized header")
    after = parse(source)
    assert replace(before, agf_metadata=None) == replace(after, agf_metadata=None)
    assert catalog_index(before, 0.55) == catalog_index(after, 0.55)
    assert (after.agf_metadata.cte2, after.agf_metadata.density, after.agf_metadata.dpgf) == (None,) * 3
    assert "Unavailable" in after.agf_metadata.provenance["dpgf"]


def test_polynomial_air_coverage_retains_existing_cap(source):
    if source[0] != "CDGM":
        pytest.skip("CDGM polynomial only")
    for column in ("AB", "AD", "AF", "AC", "AE", "AG"):
        put(source[1], source[4], column, None)
    for i, column in enumerate(("AH", "AI", "AJ", "AK", "AL", "AM")):
        put(source[1], source[4], column, 2.25 if i == 0 else 0)
    glass = parse(source)
    assert glass.dispersion_model == "polynomial"
    assert glass.agf_metadata.wavelength_range_air == (0.4, 1.014)
    assert glass.wavelength_range[1] == air_to_vacuum(1.014)


@pytest.mark.parametrize("filename,adapter,expected", [
    ("CDGM202609.xlsx", cdgm, (11.4e-6, 14.3e-6, 4.11, 0.0229)),
    ("schott-optical-glass-overview-excel-format-en 202501113.xlsx", schott, (7.4e-6, 8.5e-6, 5.41, 0.0033)),
])
def test_real_first_sheet_metadata(filename, adapter, expected):
    path = ROOT / filename
    if not path.exists():
        pytest.skip("External licensed catalog workbook not present")
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))
        glasses, _ = adapter.parse(rows, path, sheet.title, "test")
    finally:
        book.close()
    metadata = glasses[0].agf_metadata
    assert (metadata.cte1, metadata.cte2, metadata.density, metadata.dpgf) == pytest.approx(expected)
    assert all(g.agf_metadata.cte1 == g.cte for g in glasses)
