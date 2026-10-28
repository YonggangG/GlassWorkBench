"""Read-only first-sheet import of the two supported official workbook schemas."""
from hashlib import sha256
from pathlib import Path
import re
from openpyxl import load_workbook
from glass_workbench.contracts import Catalog
from . import cdgm, schott


def discover_catalogs(directory) -> list[Path]:
    """Nonrecursive recognized filenames only; legacy derived workbooks excluded."""
    directory = Path(directory)
    if not directory.is_dir():
        return []
    patterns = (r"CDGM\d{6}\.xlsx",r"schott-optical-glass-overview-excel-format-en(?:[ _-].*)?\.xlsx")
    return sorted((p for p in directory.iterdir() if p.is_file()
                   and any(re.fullmatch(pattern,p.name,re.IGNORECASE) for pattern in patterns)),
                  key=lambda p:p.name.casefold())


def load_catalog(path) -> Catalog:
    """Identify schema by FIRST worksheet headers, never fallback to later sheets."""
    path = Path(path).resolve()
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Only official .xlsx catalog workbooks are supported")
    if path.name.casefold() == "glasses_1550nm.xlsx":
        raise ValueError("Legacy derived Glasses_1550nm.xlsx is not an official catalog")
    digest = sha256(path.read_bytes()).hexdigest()
    workbook = load_workbook(path,read_only=True,data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows or not rows[0]:
            raise ValueError("Empty first worksheet")
        first = str(rows[0][0]).strip()
        if first == "Glass" and sheet.title == "optical glass":
            supplier,adapter = "CDGM",cdgm
        elif first == "SCHOTT, Optical Glass" and sheet.title == "Preferred glasses":
            supplier,adapter = "SCHOTT",schott
        else:
            raise ValueError(f"Unsupported first-sheet catalog schema: {sheet.title!r}")
        glasses,notices = adapter.parse(rows,path,sheet.title,digest)
        from .audit import audit
        glasses = audit(glasses,rows)
        for flag in ("DISPERSION_CHECKPOINT_MISMATCH","THERMAL_TABLE_MISMATCH"):
            count = sum(any(flag in w for w in g.warnings) for g in glasses)
            if count:
                notices.append(f"{count} glasses flagged {flag}; relevant calculations excluded by default")
        if not glasses:
            raise ValueError("No valid glass records in first worksheet")
        ids = [g.uid for g in glasses]
        if len(set(ids)) != len(ids):
            raise ValueError("Duplicate supplier-qualified glass identifiers in catalog")
        return Catalog(supplier,str(path),sheet.title,digest,glasses,notices)
    finally:
        workbook.close()
